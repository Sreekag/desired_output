from docx import Document

def convert_docx_to_json(file_path: str) -> dict:
    doc = Document(file_path)
    result = {"paragraphs": [], "tables": []}
    # Extract paragraphs
    result["paragraphs"] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # Extract tables
    for table in doc.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        table_data = []
        for row in table.rows[1:]:
            row_data = {headers[i]: cell.text.strip() for i, cell in enumerate(row.cells)}
            table_data.append(row_data)
        result["tables"].append(table_data)
    return result
def convert_docx_to_html(file_path: str) -> str:
    doc = Document(file_path)
    html_content = "<html><body>"
    for p in doc.paragraphs:
        if p.text.strip():
            html_content += f"<p>{p.text.strip()}</p>"
    for table in doc.tables:
        html_content += "<table border='1'>"
        html_content += "<thead><tr>"
        for cell in table.rows[0].cells:
            html_content += f"<th>{cell.text.strip()}</th>"
        html_content += "</tr></thead>"
        html_content += "<tbody>"
        for row in table.rows[1:]:
            html_content += "<tr>"
            for cell in row.cells:
                html_content += f"<td>{cell.text.strip()}</td>"
            html_content += "</tr>"
        html_content += "</tbody>"
        
        html_content += "</table><br>"
    
    html_content += "</body></html>"
    
    return html_content
